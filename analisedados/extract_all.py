"""Extract content from ALL files in analisedados directory."""
import os
import sys

# --- DOCX ---
def extract_docx(filepath):
    from docx import Document
    doc = Document(filepath)
    lines = []
    for para in doc.paragraphs:
        lines.append(para.text)
    # Also extract tables
    for i, table in enumerate(doc.tables):
        lines.append(f"\n--- TABELA {i+1} ---")
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)

# --- XLSM/XLSX ---
def extract_excel(filepath):
    from openpyxl import load_workbook
    wb = load_workbook(filepath, data_only=True, read_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"\n=== SHEET: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            vals = [str(c) if c is not None else "" for c in row]
            lines.append(" | ".join(vals))
    wb.close()
    return "\n".join(lines)

# --- PDF ---
def extract_pdf(filepath):
    from PyPDF2 import PdfReader
    reader = PdfReader(filepath)
    lines = []
    for i, page in enumerate(reader.pages):
        text = page.extract_text()
        if text:
            lines.append(f"\n--- PAGINA {i+1} ---")
            lines.append(text)
    return "\n".join(lines)

# --- MAIN ---
folder = os.path.dirname(os.path.abspath(__file__))
output_dir = os.path.join(folder, "extracted")
os.makedirs(output_dir, exist_ok=True)

files = sorted(os.listdir(folder))
for fname in files:
    if fname == "extract_all.py" or fname == "extracted":
        continue
    fpath = os.path.join(folder, fname)
    if not os.path.isfile(fpath):
        continue
    
    print(f"\n{'='*80}")
    print(f"PROCESSANDO: {fname}")
    print(f"{'='*80}")
    
    ext = os.path.splitext(fname)[1].lower()
    try:
        if ext == '.docx':
            content = extract_docx(fpath)
        elif ext in ('.xlsm', '.xlsx'):
            content = extract_excel(fpath)
        elif ext == '.pdf':
            content = extract_pdf(fpath)
        elif ext == '.doc':
            content = "[FORMATO .doc antigo - não suportado diretamente]"
        else:
            content = "[FORMATO NÃO SUPORTADO]"
        
        # Save to file
        out_name = os.path.splitext(fname)[0] + ".txt"
        out_path = os.path.join(output_dir, out_name)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write(content)
        
        # Print preview (first 3000 chars)
        preview = content[:3000]
        print(preview)
        if len(content) > 3000:
            print(f"\n... [TRUNCADO - total {len(content)} chars, salvo em {out_name}]")
        print(f"\nSalvo em: {out_path}")
        
    except Exception as e:
        print(f"ERRO ao processar {fname}: {e}")
