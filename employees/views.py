"""
Views para gestao de funcionarios.
"""
import csv
import io
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.http import HttpResponse
from django.core.paginator import Paginator
from datetime import datetime

from .models import Employee, EmployeeImportLog
from .forms import EmployeeForm, EmployeeImportForm
from companies.views import require_company_admin
from audit.models import AuditLog


@login_required
@require_company_admin
def employee_list(request):
    """Lista funcionarios da empresa."""
    company = request.user.company
    
    employees = Employee.objects.filter(company=company)
    
    search = request.GET.get('search', '')
    if search:
        employees = employees.filter(nome__icontains=search)
    
    setor = request.GET.get('setor', '')
    if setor:
        employees = employees.filter(setor=setor)
    
    status = request.GET.get('status', '')
    if status:
        employees = employees.filter(status=status)
    
    setores = Employee.objects.filter(company=company).values_list('setor', flat=True).distinct()
    
    paginator = Paginator(employees, 20)
    page = request.GET.get('page', 1)
    employees = paginator.get_page(page)
    
    context = {
        'employees': employees,
        'setores': setores,
        'search': search,
        'setor_filter': setor,
        'status_filter': status,
    }
    return render(request, 'employees/employee_list.html', context)


@login_required
@require_company_admin
def employee_create(request):
    """Cria novo funcionario."""
    company = request.user.company
    
    if not company.can_add_employee():
        messages.error(request, 'Limite de funcionarios do plano atingido.')
        return redirect('employees:list')
    
    if request.method == 'POST':
        form = EmployeeForm(request.POST, company=company)
        if form.is_valid():
            employee = form.save(commit=False)
            employee.company = company
            employee.save()
            
            AuditLog.log(
                user=request.user,
                action='CREATE',
                description=f'Funcionario {employee.nome} criado',
                obj=employee,
                request=request
            )
            
            messages.success(request, f'Funcionario {employee.nome} criado com sucesso!')
            return redirect('employees:list')
    else:
        form = EmployeeForm(company=company)
    
    return render(request, 'employees/employee_form.html', {'form': form, 'action': 'Novo'})


@login_required
@require_company_admin
def employee_edit(request, pk):
    """Edita funcionario existente."""
    company = request.user.company
    employee = get_object_or_404(Employee, pk=pk, company=company)
    
    if request.method == 'POST':
        form = EmployeeForm(request.POST, instance=employee, company=company)
        if form.is_valid():
            form.save()
            
            AuditLog.log(
                user=request.user,
                action='UPDATE',
                description=f'Funcionario {employee.nome} atualizado',
                obj=employee,
                request=request
            )
            
            messages.success(request, f'Funcionario {employee.nome} atualizado!')
            return redirect('employees:list')
    else:
        form = EmployeeForm(instance=employee, company=company)
    
    return render(request, 'employees/employee_form.html', {
        'form': form,
        'action': 'Editar',
        'employee': employee
    })


@login_required
@require_company_admin
def employee_deactivate(request, pk):
    """Desativa um funcionario."""
    company = request.user.company
    employee = get_object_or_404(Employee, pk=pk, company=company)
    
    employee.deactivate()
    
    AuditLog.log(
        user=request.user,
        action='UPDATE',
        description=f'Funcionario {employee.nome} desativado',
        obj=employee,
        request=request
    )
    
    messages.success(request, f'Funcionario {employee.nome} desativado.')
    return redirect('employees:list')


@login_required
@require_company_admin
def employee_activate(request, pk):
    """Ativa um funcionario."""
    company = request.user.company
    employee = get_object_or_404(Employee, pk=pk, company=company)
    
    employee.activate()
    
    AuditLog.log(
        user=request.user,
        action='UPDATE',
        description=f'Funcionario {employee.nome} ativado',
        obj=employee,
        request=request
    )
    
    messages.success(request, f'Funcionario {employee.nome} ativado com sucesso.')
    return redirect('employees:list')


@login_required
@require_company_admin
def employee_import(request):
    """Importa funcionarios via CSV."""
    company = request.user.company
    
    if request.method == 'POST':
        form = EmployeeImportForm(request.POST, request.FILES)
        if form.is_valid():
            uploaded_file = request.FILES['file']
            
            import_log = EmployeeImportLog.objects.create(
                company=company,
                file_name=uploaded_file.name,
                status='PROCESSING',
                created_by=request.user
            )
            
            try:
                import unicodedata
                def normalize_key(k):
                    if not k: return ''
                    k = str(k).lower().strip()
                    k = ''.join(c for c in unicodedata.normalize('NFD', k) if unicodedata.category(c) != 'Mn')
                    return k

                rows = []
                file_extension = uploaded_file.name.split('.')[-1].lower()
                
                if file_extension == 'xlsx':
                    import openpyxl
                    wb = openpyxl.load_workbook(uploaded_file, data_only=True)
                    sheet = wb.active
                    headers = [str(cell.value) if cell.value else '' for cell in sheet[1]]
                    for row in sheet.iter_rows(min_row=2, values_only=True):
                        if not any(row): continue
                        row_dict = dict(zip(headers, row))
                        rows.append(row_dict)
                else:
                    decoded_file = uploaded_file.read().decode('utf-8-sig')
                    delimiter = ';'
                    first_line = decoded_file.split('\n')[0] if decoded_file else ''
                    if ';' not in first_line and ',' in first_line:
                        delimiter = ','
                        
                    io_string = io.StringIO(decoded_file)
                    reader = csv.DictReader(io_string, delimiter=delimiter)
                    for row in reader:
                        rows.append(row)
                
                success_count = 0
                error_count = 0
                errors = []
                total_rows = 0
                
                for row_num, row in enumerate(rows, start=2):
                    total_rows += 1
                    try:
                        # Normalize row keys for flexible matching
                        row_norm = {normalize_key(k): v for k, v in row.items() if k is not None}
                        
                        def get_val(keys):
                            for k in keys:
                                norm_k = normalize_key(k)
                                if row_norm.get(norm_k): return row_norm[norm_k].strip()
                            return ''

                        nome = get_val(['nome completo', 'nome', 'funcionario', 'colaborador'])
                        if not nome:
                            raise ValueError("Nome do funcionario e obrigatorio")
                            
                        email = get_val(['e-mail corporativo', 'email', 'e-mail', 'email corporativo']).lower()
                        telefone = get_val(['telefone', 'celular', 'contato', 'tel'])
                        cpf = get_val(['cpf', 'documento'])
                        cpf = ''.join(filter(str.isdigit, cpf))
                        setor = get_val(['departamento', 'setor', 'area'])
                        cargo = get_val(['cargo/funcao', 'cargo', 'funcao'])
                        centro_de_custo = get_val(['centro de custo', 'cc'])
                        matricula = get_val(['matricula', 'registro', 're'])
                        
                        raw_admissao = get_val(['data de admissao', 'data admissao', 'admissao'])
                        raw_nascimento = get_val(['data de nascimento', 'data nascimento', 'nascimento'])
                        raw_demissao = get_val(['data de demissao', 'data demissao', 'demissao'])
                        
                        def parse_date(date_val):
                            if not date_val: return None
                            if isinstance(date_val, datetime): return date_val.date()
                            if hasattr(date_val, 'date'): return date_val.date()
                            
                            date_str = str(date_val).strip().split(' ')[0]
                            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%y', '%Y/%m/%d'):
                                try:
                                    return datetime.strptime(date_str, fmt).date()
                                except ValueError:
                                    continue
                            return None

                        data_admissao = parse_date(raw_admissao) or datetime.now().date()
                        data_nascimento = parse_date(raw_nascimento)
                        data_demissao = parse_date(raw_demissao)
                        
                        # Superior Imediato (lookup by email or name if possible)
                        gestor_val = get_val(['superior imediato', 'gestor', 'lider', 'chefe'])
                        gestor = None
                        if gestor_val:
                            gestor = Employee.objects.filter(company=company, email__iexact=gestor_val).first()
                            if not gestor:
                                gestor = Employee.objects.filter(company=company, nome__icontains=gestor_val).first()

                        status_val = get_val(['status', 'situacao', 'estado']).upper()
                        status = 'ACTIVE'
                        if 'INATIVO' in status_val or 'OFF' in status_val or 'TERMINATED' in status_val or 'DESLIGADO' in status_val:
                            status = 'TERMINATED'
                        elif 'AFASTADO' in status_val:
                            status = 'ON_LEAVE'
                            
                        turno_val = get_val(['turno', 'horario']).upper()
                        turno = 'FULL'
                        if 'MANHA' in turno_val: turno = 'MORNING'
                        elif 'TARDE' in turno_val: turno = 'AFTERNOON'
                        elif 'NOITE' in turno_val: turno = 'NIGHT'
                        elif 'REVEZAMENTO' in turno_val: turno = 'ROTATING'

                        employee, created = Employee.objects.update_or_create(
                            company=company,
                            email=email,
                            defaults={
                                'nome': nome,
                                'cpf': cpf,
                                'telefone': telefone,
                                'setor': setor,
                                'cargo': cargo,
                                'centro_de_custo': centro_de_custo,
                                'turno': turno,
                                'data_admissao': data_admissao,
                                'data_nascimento': data_nascimento,
                                'data_demissao': data_demissao,
                                'matricula': matricula,
                                'gestor': gestor,
                                'status': status
                            }
                        )
                        success_count += 1
                        
                    except Exception as e:
                        error_count += 1
                        errors.append({
                            'row': row_num,
                            'data': dict(row),
                            'error': str(e)
                        })
                
                import_log.total_rows = total_rows
                import_log.success_count = success_count
                import_log.error_count = error_count
                import_log.errors = errors
                import_log.status = 'COMPLETED' if error_count == 0 else 'COMPLETED'
                import_log.save()
                
                AuditLog.log(
                    user=request.user,
                    action='IMPORT',
                    description=f'Importacao de {success_count} funcionarios via CSV',
                    obj=import_log,
                    request=request
                )
                
                messages.success(
                    request,
                    f'Importacao concluida: {success_count} importados, {error_count} erros.'
                )
                
            except Exception as e:
                import_log.status = 'FAILED'
                import_log.errors = [{'error': str(e)}]
                import_log.save()
                messages.error(request, f'Erro na importacao: {str(e)}')
            
            return redirect('employees:list')
    else:
        form = EmployeeImportForm()
    
    return render(request, 'employees/employee_import.html', {'form': form})


@login_required
@require_company_admin
def employee_export_template(request):
    """Exporta template Excel para importacao."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Funcionarios"

    headers = [
        'Nome Completo', 'E-mail Corporativo', 'Telefone', 'CPF', 'Departamento', 'Cargo/Função',
        'Centro de Custo', 'Superior Imediato', 'Data de Nascimento', 'Data de Admissão', 'Matrícula', 'Turno', 'Status'
    ]
    
    ws.append(headers)
    
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
        ws.column_dimensions[get_column_letter(col_num)].width = len(header) + 5

    ws.append([
        'João da Silva', 'joao@empresa.com.br', '(11) 99999-9999', '12345678901', 'TI', 'Analista de Sistemas',
        'CC-001', 'gestor@empresa.com.br', '15/05/1985', '01/01/2024', '001', 'INTEGRAL', 'ATIVO'
    ])

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="modelo_funcionarios.xlsx"'
    wb.save(response)
    
    return response


@login_required
@require_company_admin
def employee_create_user(request, pk):
    """Cria conta de usuario para funcionario."""
    company = request.user.company
    employee = get_object_or_404(Employee, pk=pk, company=company)
    
    if employee.user:
        messages.warning(request, 'Este funcionario ja possui uma conta de usuario.')
        return redirect('employees:list')
    
    user = employee.create_user_account()
    
    AuditLog.log(
        user=request.user,
        action='CREATE',
        description=f'Conta de usuario criada para {employee.nome}',
        obj=user,
        request=request
    )
    
    messages.success(
        request,
        f'Conta criada para {employee.nome}. A senha de acesso padrao é o CPF do funcionario (apenas numeros).'
    )
    return redirect('employees:list')
